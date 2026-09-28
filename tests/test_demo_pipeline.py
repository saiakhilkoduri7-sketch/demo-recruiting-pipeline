from app.demo_pipeline import load_resume_candidates, run_demo_pipeline


def test_load_resume_candidates_reads_docx_files(tmp_path):
    docx_file = tmp_path / "charlie.docx"

    try:
        from docx import Document
    except Exception as exc:
        raise AssertionError("python-docx is required for this test") from exc

    document = Document()
    document.add_paragraph("Charlie Candidate")
    document.add_paragraph("Python, Django, SQL")
    document.save(docx_file)

    candidates = load_resume_candidates(tmp_path)

    assert len(candidates) == 1
    assert candidates[0].name == "charlie"
    assert "django" in candidates[0].resume_text.lower()


def test_load_resume_candidates_reads_text_files(tmp_path):
    resume_file = tmp_path / "alice.txt"
    resume_file.write_text("Alice Applicant\nPython FastAPI PostgreSQL", encoding="utf-8")

    candidates = load_resume_candidates(resume_file.parent)

    assert len(candidates) == 1
    assert candidates[0].name == "alice"
    assert "python" in candidates[0].resume_text.lower()


def test_run_demo_pipeline_returns_ranked_matches(tmp_path):
    resume_one = tmp_path / "alice.txt"
    resume_one.write_text("Alice Applicant\nPython FastAPI PostgreSQL", encoding="utf-8")

    resume_two = tmp_path / "bob.txt"
    resume_two.write_text("Bob Candidate\nJava Spring Boot AWS", encoding="utf-8")

    results = run_demo_pipeline(
        resume_folder=str(tmp_path),
        job_title="Senior Python Engineer",
        job_description="Need Python, FastAPI, PostgreSQL, cloud deployment experience.",
    )

    assert results["agent1"]["unique_candidates"] == 2
    assert results["agent1"]["duplicates"] == 0
    assert results["agent2"][0]["candidate_name"] == "alice"
    assert results["agent2"][0]["score"] > results["agent2"][1]["score"]
